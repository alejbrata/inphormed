import os, glob, hashlib
from datetime import datetime
from uuid import uuid4

from dotenv import load_dotenv
from qdrant_client import QdrantClient
from qdrant_client.http import models as qm
from sentence_transformers import SentenceTransformer

# ----------------- Config -----------------
BASE_DIR = os.path.dirname(os.path.dirname(__file__))
load_dotenv(os.path.join(BASE_DIR, ".env"))

QDRANT_URL = os.getenv("QDRANT_URL")
QDRANT_API_KEY = os.getenv("QDRANT_API_KEY")
COLLECTION = os.getenv("QDRANT_COLLECTION", "biomed_evidence_v1")
EMBED_DIM = int(os.getenv("EMBED_DIM", "384"))

INPUT_DIR = os.path.join(BASE_DIR, "uploads", "claims")   # pon aquí tus DOCX/PPTX
MODEL_NAME = "sentence-transformers/paraphrase-multilingual-MiniLM-L12-v2"

# ----------------- Utilidades -----------------
def sha256(text: str) -> str:
    return hashlib.sha256(text.encode("utf-8")).hexdigest()

def norm(s: str) -> str:
    return " ".join(s.split()).strip()

# ----------------- Extractores -----------------
def extract_claims_docx(path: str):
    """Devuelve lista de dicts: {'text', 'where'} desde un .docx."""
    from docx import Document
    doc = Document(path)
    claims = []
    for i, p in enumerate(doc.paragraphs, start=1):
        t = norm(p.text)
        if not t:
            continue
        # Heurística simple: capturamos bullets/num y líneas no vacías
        if p.style and ("List" in p.style.name or "Bullet" in p.style.name):
            claims.append({"text": t, "where": f"docx:paragraph:{i}"})
        elif len(t) > 15:
            claims.append({"text": t, "where": f"docx:paragraph:{i}"})
    return claims

def extract_claims_pptx(path: str):
    """Devuelve lista de dicts: {'text','where'} desde un .pptx."""
    from pptx import Presentation
    prs = Presentation(path)
    claims = []
    for si, slide in enumerate(prs.slides, start=1):
        for sh in slide.shapes:
            if not hasattr(sh, "text"):
                continue
            t = norm(sh.text)
            if not t:
                continue
            # Heurística: títulos y cuadros de texto con 1–2 frases
            if len(t) >= 10:
                claims.append({"text": t, "where": f"pptx:slide:{si}"})
    return claims

def extract_claims_any(path: str):
    ext = os.path.splitext(path)[1].lower()
    if ext == ".docx":
        return extract_claims_docx(path)
    if ext == ".pptx":
        return extract_claims_pptx(path)
    return []

# ----------------- Qdrant / Embeddings -----------------
def ensure_collection(client: QdrantClient):
    if not client.collection_exists(COLLECTION):
        client.create_collection(
            collection_name=COLLECTION,
            vectors_config=qm.VectorParams(size=EMBED_DIM, distance=qm.Distance.COSINE),
        )

def embed_texts(model, texts):
    vecs = model.encode(texts, convert_to_numpy=True, normalize_embeddings=True)
    assert vecs.shape[1] == EMBED_DIM, f"Esperaba {EMBED_DIM}, obtuve {vecs.shape[1]}"
    return vecs

# ----------------- Main -----------------
def main():
    assert QDRANT_URL, "Falta QDRANT_URL"
    client = QdrantClient(url=QDRANT_URL, api_key=QDRANT_API_KEY)
    ensure_collection(client)

    files = glob.glob(os.path.join(INPUT_DIR, "*.docx")) + glob.glob(os.path.join(INPUT_DIR, "*.pptx"))
    if not files:
        print(f"⚠️ No hay archivos en {INPUT_DIR}")
        return

    model = SentenceTransformer(MODEL_NAME)
    now = datetime.utcnow().isoformat() + "Z"

    total_points = 0
    for path in files:
        fname = os.path.basename(path)
        source = "docx" if fname.lower().endswith(".docx") else "pptx"
        # doc_id: estable por archivo (checksum del archivo)
        with open(path, "rb") as f:
            file_hash = hashlib.sha256(f.read()).hexdigest()[:16]
        doc_id = f"{source}:{fname}:{file_hash}"

        claims = extract_claims_any(path)
        if not claims:
            print(f"— {fname}: 0 claims detectados")
            continue

        texts = [c["text"] for c in claims]
        vecs = embed_texts(model, texts)

        points = []
        for c, v in zip(claims, vecs):
            text = c["text"]
            cid = sha256(text)[:16]  # id estable por texto
            payload = {
                "doc_id": doc_id,
                "claim_id": cid,
                "source": source,
                "file_name": fname,
                "where": c["where"],      # slide/párrafo
                "text": text,
                "year": datetime.utcnow().year,  # si aplica
                "ingested_at": now,
                "kind": "claim",          # para filtrar claims vs papers
            }
            points.append(qm.PointStruct(id=str(uuid4()), vector=v.tolist(), payload=payload))

        client.upsert(collection_name=COLLECTION, points=points)
        total_points += len(points)
        print(f"✅ {fname}: {len(points)} claims indexados → doc_id={doc_id}")

    print(f"🎯 Total upserts: {total_points}")

if __name__ == "__main__":
    main()
