# app/api/claims.py
from __future__ import annotations

import os
from io import BytesIO
from typing import Any, Dict, List, Optional

from fastapi import APIRouter, UploadFile, File, HTTPException, Query
from pydantic import BaseModel

from app.config.settings import settings
from app.audit.auditor import Auditor
from app.indexers.vector_indexer import VectorIndexer
from app.utils.query_rewriter import QueryRewriter
from app.agents.claims.validator import ClaimsValidator
from app.workflow.coordinator import WorkflowCoordinator

# 👉 Usa el registro en INGLÉS (si tu carpeta es app/agents/sources)
from app.agents.sources.registry import get_default_sources as get_default_fuentes
# Si dejaste alias en español, podrías usar:
# from app.agents.fuentes.registry import get_default_fuentes

router = APIRouter(prefix="/api/claims", tags=["claims"])

# ─────────────────────────────────────────────────────────────
# Auditor seguro (acepta distintas firmas o hace fallback)
# ─────────────────────────────────────────────────────────────
_log_path = getattr(settings, "AUDIT_LOG_PATH", "logs/audit.log")
os.makedirs(os.path.dirname(_log_path), exist_ok=True)

def _build_auditor():
    try:
        # Preferimos keyword si tu clase lo soporta
        return Auditor(log_path=_log_path)
    except TypeError:
        try:
            # Algunas implementaciones sólo aceptan posicional
            return Auditor(_log_path)
        except TypeError:
            try:
                # Otras no aceptan path alguno
                return Auditor()
            except Exception:
                # Fallback nulo, pero con misma interfaz
                class _NullAuditor:
                    def log_event(self, *args, **kwargs): ...
                return _NullAuditor()

auditor = _build_auditor()

# ─────────────────────────────────────────────────────────────
# Singletons de orquestación
# ─────────────────────────────────────────────────────────────
indexer = VectorIndexer(
    qdrant_url=settings.QDRANT_URL,
    collection=settings.QDRANT_COLLECTION,
    api_key=settings.QDRANT_API_KEY or None,
    embedding_model=settings.EMBEDDING_MODEL,
)
rewriter = QueryRewriter()
validator = ClaimsValidator()
coordinator = WorkflowCoordinator(
    query_rewriter=rewriter,
    validator=validator,
    citations_limit=getattr(settings, "CITATIONS_LIMIT", 3),
    min_hits=getattr(settings, "MIN_HITS", 1),
    indexer=indexer,
    auditor=auditor,
)
agentes_fuente = get_default_fuentes(indexer=indexer, auditor=auditor)

# ─────────────────────────────────────────────────────────────
# Modelos de entrada
# ─────────────────────────────────────────────────────────────
class ValidateTextPayload(BaseModel):
    text: str
    topk: Optional[int] = 8
    thr_green: Optional[float] = 0.82
    thr_yellow: Optional[float] = 0.70
    require_llm: Optional[bool] = True

# ─────────────────────────────────────────────────────────────
# Helpers de formato de respuesta
# ─────────────────────────────────────────────────────────────
def _label_to_status(lbl: str) -> str:
    l = (lbl or "").upper()
    if l in ("GREEN", "VERDE"): return "green"
    if l in ("YELLOW", "AMBER", "AMARILLO"): return "yellow"
    return "red"

def _citations_to_list(citations: Any) -> List[Dict[str, Any]]:
    out: List[Dict[str, Any]] = []
    for c in (citations or []):
        out.append({
            "source": getattr(c, "source", None),
            "url": getattr(c, "url", None),
            "title": getattr(c, "title", None),
            "span_start": getattr(c, "span_start", None),
            "span_end": getattr(c, "span_end", None),
            "doc_hash": getattr(c, "doc_hash", None),
            "score": getattr(c, "score", None),
        })
    return out

def _best_score(citations: List[Dict[str, Any]]) -> float:
    if not citations:
        return 0.0
    scs = [c.get("score") for c in citations if isinstance(c.get("score"), (int, float))]
    return float(max(scs)) if scs else 0.0

# ─────────────────────────────────────────────────────────────
# Endpoints
# ─────────────────────────────────────────────────────────────
@router.post("/validate-text")
def validate_text(payload: ValidateTextPayload) -> Dict[str, Any]:
    """
    Valida UN claim:
      - QueryRewriter (LLM) → RAG
      - Si no hay evidencia: agentes (PubMed/CTGov/EMA) → ingesta → reintento
      - ClaimsValidator (LLM) + Compliance
    """
    claim_text = (payload.text or "").strip()
    if not claim_text:
        raise HTTPException(status_code=400, detail="Campo 'text' vacío.")

    out = coordinator.resolver_claim(
        claim=claim_text,
        intake_id="text",
        claim_id="text-0",
        agentes_fuente=agentes_fuente,
        generated_text=claim_text,
        meta={},
        require_llm=bool(payload.require_llm),
    )

    val = out["resultado"]
    comp = out["compliance"]
    models = out.get("models") or {}

    label = getattr(val, "label", "RED")
    top_url = getattr(val, "top_url", None)
    citations = _citations_to_list(getattr(val, "citations", []))
    status = _label_to_status(label)

    return {
        "where": "text",
        "text": claim_text,
        "claim_id": "text-0",
        "status": status,
        "best_score": _best_score(citations),
        "top_url": top_url,
        "citations": citations,   # nuevo
        "hits": citations,        # compat con UI actual
        "compliance": {
            "score": getattr(comp, "score", None),
            "passed": getattr(comp, "passed", None),
            "issues": getattr(comp, "issues", None),
        },
        "models": models,
    }

@router.post("/validate-ppt")
async def validate_ppt(
    file: UploadFile = File(...),
    topk: int = Query(8),
    thr_green: float = Query(0.82),
    thr_yellow: float = Query(0.70),
    require_llm: bool = Query(True),
    render_ppt: bool = Query(False),
) -> Dict[str, Any]:
    """
    Valida TODOS los párrafos de un PPTX (cada párrafo ≈ 1 claim).
    Devuelve estructura compatible con tu Streamlit.
    """
    if file.content_type not in (
        "application/vnd.openxmlformats-officedocument.presentationml.presentation",
        "application/octet-stream",
    ):
        raise HTTPException(status_code=400, detail="Sube un .pptx válido.")

    # Leer PPTX de manera segura (BytesIO)
    from pptx import Presentation
    content = await file.read()
    try:
        prs = Presentation(BytesIO(content))
    except Exception:
        raise HTTPException(status_code=400, detail="PPTX corrupto o no válido.")

    results: List[Dict[str, Any]] = []
    total_claims = 0

    for si, slide in enumerate(prs.slides, start=1):
        for shi, shape in enumerate(slide.shapes, start=1):
            if not getattr(shape, "has_text_frame", False):
                continue
            for pi, para in enumerate(shape.text_frame.paragraphs, start=1):
                txt = (para.text or "").strip()
                if not txt:
                    continue
                total_claims += 1
                claim_id = f"slide{si}-shape{shi}-p{pi}"

                out = coordinator.resolver_claim(
                    claim=txt,
                    intake_id=f"ppt::{file.filename}",
                    claim_id=claim_id,
                    agentes_fuente=agentes_fuente,
                    generated_text=txt,
                    meta={},
                    require_llm=bool(require_llm),
                )

                val = out["resultado"]
                comp = out["compliance"]
                models = out.get("models") or {}
                label = getattr(val, "label", "RED")
                top_url = getattr(val, "top_url", None)
                citations = _citations_to_list(getattr(val, "citations", []))
                status = _label_to_status(label)

                results.append({
                    "where": f"slide:{si}",
                    "text": txt,
                    "claim_id": claim_id,
                    "status": status,
                    "best_score": _best_score(citations),
                    "top_url": top_url,
                    "citations": citations,   # nuevo
                    "hits": citations,        # compat con UI
                    "compliance": {
                        "score": getattr(comp, "score", None),
                        "passed": getattr(comp, "passed", None),
                        "issues": getattr(comp, "issues", None),
                    },
                    "models": models,
                })

    return {
        "file_name": file.filename,
        "total_claims": total_claims,
        "results": results,
        # "ppt_b64": ..., "ppt_name": ...  # si activas el renderizado del PPT coloreado
    }
