# app/services/claim_extractor.py
from __future__ import annotations
from io import BytesIO
from typing import List, Dict
from datetime import datetime
import hashlib

__all__ = ["extract_claims_from_pptx"]

def _norm(s: str) -> str:
    return " ".join((s or "").split()).strip()

def _sha16(text: str) -> str:
    return hashlib.sha256(text.encode("utf-8")).hexdigest()[:16]

def _file_sha16(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()[:16]

def extract_claims_from_pptx(file_bytes: bytes, file_name: str) -> Dict:
    """
    Extrae 'claims' de un PPTX en memoria usando heurística simple
    (cajas de texto/títulos con ≥10 caracteres).
    Devuelve:
    {
      "doc_id": str,
      "file_name": str,
      "claims": [ { "text","where","claim_id","source","file_name","doc_id" }, ... ],
      "ingested_at": ISO8601
    }
    """
    try:
        from pptx import Presentation
    except Exception as e:
        raise RuntimeError("Falta dependencia 'python-pptx' (pip install python-pptx)") from e

    bio = BytesIO(file_bytes)
    prs = Presentation(bio)
    file_hash = _file_sha16(file_bytes)
    doc_id = f"pptx:{file_name}:{file_hash}"

    out: List[Dict] = []
    for si, slide in enumerate(prs.slides, start=1):
        for sh in slide.shapes:
            # Prioriza text_frame (cuadros de texto, títulos)
            if hasattr(sh, "text_frame") and sh.text_frame:
                t = _norm(sh.text_frame.text or "")
            elif hasattr(sh, "text"):
                t = _norm(sh.text)
            else:
                t = ""
            if not t or len(t) < 10:
                continue

            where = f"slide:{si}"
            claim_id = _sha16(t)
            out.append({
                "text": t,
                "where": where,
                "claim_id": claim_id,
                "source": "pptx",
                "file_name": file_name,
                "doc_id": doc_id,
            })

    return {
        "doc_id": doc_id,
        "file_name": file_name,
        "claims": out,
        "ingested_at": datetime.utcnow().isoformat() + "Z",
    }
