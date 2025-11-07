# -*- coding: utf-8 -*-
"""
Extractor robusto de claims desde PPTX sin imponer formato al origen.
Diferencia frases-claim de bloques de datos/soporte en la misma slide.

Dependencias:
  - python-pptx

Uso:
  from app.claims.pptx_parser import extract_claims_from_pptx
  claims = extract_claims_from_pptx("claims_ppt_01.pptx", max_claims_per_slide=1)

Devuelve una lista de dicts:
  {
    "slide_index": int,
    "text": str,
    "score": float,          # 0..1 confianza "esto es claim"
    "cid": str|None,         # si podemos inferir algún id
    "debug": { ... }         # señales usadas para puntuar
  }
"""

from __future__ import annotations
from dataclasses import dataclass, asdict
from typing import List, Dict, Any, Tuple
from pptx import Presentation
from pptx.enum.shapes import PP_PLACEHOLDER
from difflib import SequenceMatcher
import re
import hashlib

# --- Léxicos y patrones (ES + EN) -------------------------------------------

CLAIM_LEX_ES = {
    # verbos y patrones típicos de afirmación/aseveración clínica
    "reduce", "redujo", "reducción", "disminuye", "disminuyó",
    "mejora", "mejoró", "incrementa", "incrementó", "aumenta", "aumentó",
    "superior", "no inferior", "no-inferior", "equivalente",
    "demuestra", "demostró", "evidencia", "asociado", "asocia", "logra", "alcanza",
    "eficaz", "efectivo", "seguro", "superioridad", "no inferioridad",
    "reduce el riesgo", "menor riesgo", "mayor supervivencia", "mejor supervivencia",
}
CLAIM_LEX_EN = {
    "reduces", "reduced", "improves", "improved", "increases", "increased",
    "superior", "noninferior", "non-inferior", "equivalent",
    "demonstrates", "demonstrated", "evidence", "associated", "achieves",
    "effective", "efficacious", "safe", "superiority", "noninferiority",
    "reduces risk", "lower risk", "higher survival", "improved survival",
}
# encabezados o palabras que suelen marcar "datos" (restan peso de claim)
DATA_TOKENS_ES = {
    "datos", "detalle", "detalles", "metodolog", "método", "materiales",
    "resultados", "análisis", "anexo", "tabla", "tablas", "figura", "figuras",
    "población", "grupo control", "ensayo", "estudio", "variable",
}
DATA_TOKENS_EN = {
    "data", "details", "methods", "methodology", "materials",
    "results", "appendix", "table", "tables", "figure", "figures",
    "population", "control group", "trial", "study", "variable",
}

# Patrones que indican soporte/estadísticos, no claim:
PAT_REF = re.compile(r"(pmid|doi|https?://|et al\.|[12][09]\d{2})", re.I)
PAT_P = re.compile(r"\bp\s*([<=>]|=)\s*0?\.\d+\b", re.I)
PAT_CI = re.compile(r"(ic\s*95%|95%\s*ci|ci\s*95%)", re.I)
PAT_N = re.compile(r"\b[nsN]\s*=\s*\d+|\bn=\d+", re.I)
PAT_NUM = re.compile(r"\d{1,3}(\.\d+)?\s*%")
PAT_METRIC = re.compile(r"\b(hr|or|rr|sd|se|i2|i²|md)\b", re.I)

# Prefijos útiles
PAT_CLAIM_PREFIX = re.compile(r"^\s*(claim|conclusión|conclusion)[:\s#-]", re.I)

# --- Utilidades --------------------------------------------------------------

def _norm(txt: str) -> str:
    t = txt.strip().lower()
    t = re.sub(r"\s+", " ", t)
    return t

def _similar(a: str, b: str) -> float:
    return SequenceMatcher(None, a, b).ratio()

def _split_candidates(text: str) -> List[str]:
    """
    Divide una caja de texto en posibles frases-claim:
    - separa por saltos y por puntuación fuerte.
    - limpia bullets.
    """
    if not text:
        return []
    t = text.replace("•", "\n").replace("▪", "\n").replace("·", "\n")
    # Rompemos por nueva línea y por punto + espacio, pero mantenemos abreviaturas simples.
    raw = re.split(r"\n+|(?<=[^A-ZÁÉÍÓÚA-Z0-9])\.\s+", t)
    cand = []
    for chunk in raw:
        c = re.sub(r"\s+", " ", chunk).strip()
        c = re.sub(PAT_CLAIM_PREFIX, "", c)  # quita "Claim:" si está
        # descartes rápidos
        if not c:
            continue
        if len(c) < 20:         # muy corto para ser claim útil
            continue
        if len(c) > 300:        # demasiado largo: probablemente bloque de datos
            continue
        cand.append(c)
    return cand

def _score_line(line: str, slide_title: str, shape_name: str, is_title: bool, is_body: bool) -> Tuple[float, Dict[str, Any]]:
    """
    Puntúa 0..1 "parece un claim". Heurística sumatoria con límites.
    """
    dbg: Dict[str, Any] = {}
    score = 0.0

    ln = _norm(line)
    title = _norm(slide_title or "")
    name = (shape_name or "")

    # 1) léxico pro-claim
    has_es = any(tok in ln for tok in CLAIM_LEX_ES)
    has_en = any(tok in ln for tok in CLAIM_LEX_EN)
    if has_es: score += 0.35
    if has_en: score += 0.30
    dbg["has_claim_lex_es"] = has_es
    dbg["has_claim_lex_en"] = has_en

    # 2) señales anti-claim (datos)
    ref = bool(PAT_REF.search(line))
    pval = bool(PAT_P.search(line))
    ci = bool(PAT_CI.search(line))
    nval = bool(PAT_N.search(line))
    pct = bool(PAT_NUM.search(line))
    metric = bool(PAT_METRIC.search(line))
    anti_hits = sum([ref, pval, ci, nval, pct, metric])
    score -= anti_hits * 0.18
    dbg.update({"has_ref": ref, "has_p": pval, "has_ci": ci, "has_n": nval, "has_pct": pct, "has_metric": metric})

    # 3) contexto del título de slide (si huele a "datos", restamos)
    data_title_hit = any(tok in title for tok in DATA_TOKENS_ES | DATA_TOKENS_EN)
    if data_title_hit:
        score -= 0.25
    dbg["data_title_hit"] = data_title_hit

    # 4) forma/placeholder (título y cuerpo tienen más papeletas que shapes raras)
    if is_title:
        score += 0.08
    if is_body:
        score += 0.05
    dbg["is_title_ph"] = is_title
    dbg["is_body_ph"] = is_body

    # 5) estructura lingüística mínima de afirmación (verbo + objeto)
    verb_like = bool(re.search(r"\b(reduce|reduces|mejora|improves|aumenta|increases|demuestra|demonstrates)\b", ln))
    comparative = bool(re.search(r"\b(mayor|menor|superior|inferior|higher|lower|superior|non[- ]?inferior)\b", ln))
    if verb_like:   score += 0.12
    if comparative: score += 0.07
    dbg["verb_like"] = verb_like
    dbg["comparative"] = comparative

    # 6) castigo por URLs/PMID/DOI explícitos (ya cubierto por ref, pero reforzamos)
    if "http" in ln or "pmid" in ln or "doi" in ln:
        score -= 0.1

    # 7) penalización por paréntesis/técnicismos densos (suele ser soporte)
    paren_ratio = line.count("(") + line.count(")")
    if paren_ratio >= 2:
        score -= 0.05
    dbg["paren_count"] = paren_ratio

    # 8) longitud "natural" de claim (preferible 40-200 chars)
    L = len(line)
    if 40 <= L <= 200:
        score += 0.08
    elif L < 35:
        score -= 0.05
    else:
        score -= 0.04
    dbg["len"] = L

    # clamp 0..1
    score = max(0.0, min(1.0, score))
    return score, dbg

@dataclass
class ClaimHit:
    slide_index: int
    text: str
    score: float
    cid: str | None
    debug: Dict[str, Any]

def _dedupe(claims: List[ClaimHit], threshold: float = 0.90) -> List[ClaimHit]:
    out: List[ClaimHit] = []
    seen: set[str] = set()
    for c in sorted(claims, key=lambda x: x.score, reverse=True):
        n = _norm(c.text)
        h = hashlib.sha1(n.encode("utf-8")).hexdigest()
        if h in seen:
            continue
        if any(_similar(_norm(o.text), n) >= threshold for o in out):
            continue
        seen.add(h)
        out.append(c)
    return out

# --- Extractor principal -----------------------------------------------------

def extract_claims_from_pptx(path: str, max_claims_per_slide: int = 1, global_dedupe: bool = True) -> List[Dict[str, Any]]:
    prs = Presentation(path)
    hits: List[ClaimHit] = []

    for idx, slide in enumerate(prs.slides):
        title = ""
        try:
            title = slide.shapes.title.text or ""
        except Exception:
            title = ""

        # Tomamos solo shapes textuales, fuera tablas/grupos
        shape_candidates: List[Tuple[str, str, bool, bool]] = []
        for sh in slide.shapes:
            if getattr(sh, "has_table", False) and sh.has_table:
                continue
            if not getattr(sh, "has_text_frame", False):
                continue
            txt = (sh.text or "").strip()
            if not txt:
                continue
            is_title = False
            is_body = False
            try:
                if getattr(sh, "is_placeholder", False):
                    pht = sh.placeholder_format.type
                    is_title = (pht == PP_PLACEHOLDER.TITLE)
                    is_body = (pht == PP_PLACEHOLDER.BODY)
            except Exception:
                pass
            shape_candidates.append((txt, getattr(sh, "name", "") or "", is_title, is_body))

        # De cada shape, dividimos en frases candidatas y puntuamos
        line_scores: List[Tuple[float, str, Dict[str, Any]]] = []
        for txt, name, is_title, is_body in shape_candidates:
            for cand in _split_candidates(txt):
                sc, dbg = _score_line(cand, title, name, is_title, is_body)
                dbg.update({"shape_name": name, "slide_title": title})
                line_scores.append((sc, cand, dbg))

        if not line_scores:
            continue

        # Ordenamos por score y nos quedamos con los top-N por slide
        line_scores.sort(key=lambda x: x[0], reverse=True)
        chosen = [ls for ls in line_scores if ls[0] >= 0.50]  # umbral mínimo de claim
        chosen = chosen[:max_claims_per_slide] if max_claims_per_slide > 0 else chosen

        for sc, cand, dbg in chosen:
            hits.append(ClaimHit(
                slide_index=idx,
                text=cand.strip(),
                score=sc,
                cid=None,
                debug=dbg
            ))

    # Deduplicación global
    hits = _dedupe(hits, threshold=0.90) if global_dedupe else hits

    # Orden estable por slide y score
    hits.sort(key=lambda h: (h.slide_index, -h.score, h.text))

    return [asdict(h) for h in hits]
